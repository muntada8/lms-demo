import sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = sys.argv[1] if len(sys.argv) > 1 else "Power_Automate_Readiness_Test.docx"
BLUE = RGBColor(0x0B, 0x4F, 0x9C)
GREY = RGBColor(0x55, 0x55, 0x55)

doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(1.8)
    s.left_margin = s.right_margin = Cm(2.0)
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(10.5)
st.paragraph_format.space_after = Pt(3)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def heading(text, size=13):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = BLUE
    return p


def para(text="", bold_prefix=None, italic=False, color=None, size=None):
    p = doc.add_paragraph()
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
    r = p.add_run(text)
    r.italic = italic
    if color:
        r.font.color.rgb = color
    if size:
        r.font.size = Pt(size)
    return p


def bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(1)
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
    p.add_run(text)
    return p


def fix_widths(t, widths):
    t.autofit = False
    grid = t._tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(widths[i] * 567)))
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)


def table(header, rows, widths, header_fill="0B4F9C"):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(10)
        shade(c, header_fill)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            if isinstance(v, tuple):  # (bold, normal)
                p.add_run(v[0]).bold = True
                p.add_run(v[1])
            else:
                p.add_run(v)
            for run in p.runs:
                run.font.size = Pt(10)
    fix_widths(t, widths)
    return t


def callout(title, lines, fill):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c = t.rows[0].cells[0]
    shade(c, fill)
    c.text = ""
    p = c.paragraphs[0]
    p.add_run(title).bold = True
    for ln in lines:
        q = c.add_paragraph()
        q.paragraph_format.space_after = Pt(1)
        if isinstance(ln, tuple):
            q.add_run(ln[0]).bold = True
            q.add_run(ln[1])
        else:
            q.add_run(ln)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


# ---------- Title ----------
p = doc.add_paragraph()
r = p.add_run("PRE-TRAINING READINESS TEST")
r.bold = True
r.font.size = Pt(18)
r.font.color.rgb = BLUE
p = doc.add_paragraph()
r = p.add_run("Power Automate: Beginner to Intermediate — Is my account ready?")
r.font.size = Pt(12)
r.font.color.rgb = GREY
para("Time needed: about 15–20 minutes. Do this on your office laptop, with your office account, "
     "at least 1 week before the training. You do not need any Power Automate experience — just follow the steps.",
     italic=True)

table(["Name", "", "Department", ""], [["Office email", "", "Date tested", ""]],
      [3, 5.5, 3, 5.5], header_fill="0B4F9C")

# ---------- What you DON'T need ----------
heading("1. Good news first — what this training does NOT need")
para("Many companies block some Microsoft features for compliance. The items below are blocked in many "
     "organisations, and that is fine — the training does not use them:")
table(["Not needed", "Why"], [
    ["Office Store Add-ins (Excel / Outlook / Word add-ins)",
     "Power Automate is a separate web app (make.powerautomate.com). It does not run as an add-in."],
    ["Microsoft Copilot / Copilot in Power Automate",
     "All flows are built by hand, step by step. If you see “Copilot is not available”, just ignore it."],
    ["SharePoint", "All labs use OneDrive and Excel. SharePoint is shown by the trainer as a demo only."],
    ["Teams meetings", "We only send a Teams chat message — and if Teams is blocked, we use Outlook email instead."],
    ["Premium connectors, HTTP, custom connectors, AI Builder, desktop flows (RPA)",
     "Not covered in this course. Only standard Microsoft 365 connectors are used."],
    ["Sending to external / personal email", "All tests and labs send email only to your own office address."],
], [6.5, 10.5])

# ---------- The test ----------
heading("2. The test — build one small test flow")
para("Open Chrome or Edge and go to ", bold_prefix=None).add_run("https://make.powerautomate.com").bold = True
para("Sign in with your office email. Then do each step in order and tick the result. "
     "If any step shows an error, do not try to fix it — take a screenshot and go to the next step.")

table(["#", "What to do", "What you should see", "OK?"], [
    ["1", ("Sign in. ", "Go to make.powerautomate.com and sign in with your office email."),
     "The Power Automate home page opens. No “Start free trial”, “You don’t have a license” or “Access denied” message.", "☐"],
    ["2", ("Create a flow. ", "Click + Create → Instant cloud flow. Name it READINESS TEST. "
                             "Choose “Manually trigger a flow” → Create."),
     "The flow designer opens with a “Manually trigger a flow” box.", "☐"],
    ["3", ("Outlook. ", "Click + (New step) → search “Send an email (V2)” (Office 365 Outlook). "
                       "To: your own office email. Subject: Test OK. Body: Hello."),
     "The action is added. If asked to sign in to create a connection, sign in — it should succeed.", "☐"],
    ["4", ("OneDrive. ", "Add action “Create file” (OneDrive for Business). Folder: / (root). "
                        "File name: readiness-test.txt. File content: Test."),
     "The action is added with no connection error.", "☐"],
    ["5", ("Excel. ", "First, in OneDrive, create an Excel file named Test.xlsx; type Name in A1 and "
                     "Test in A2; select A1:A2 → Insert → Table. Then add action "
                     "“List rows present in a table” (Excel Online (Business)) and pick that file and table."),
     "You can browse to Test.xlsx and the table (Table1) appears in the drop-down.", "☐"],
    ["6", ("Approvals. ", "Add action “Start and wait for an approval”. Type: Approve/Reject – First to respond. "
                         "Title: Test approval. Assigned to: your own office email."),
     "The action is added with no error.", "☐"],
    ["7", ("Teams (optional). ", "Add action “Post message in a chat or channel” (Microsoft Teams). "
                                "Post as: Flow bot. Post in: Chat with Flow bot. Recipient: your own email. Message: Test."),
     "The action is added with no error. (If Teams is blocked, tick N/A — we will use email instead.)", "☐ / N/A"],
    ["8", ("Save. ", "Click Save (top right)."),
     "“Your flow is ready” / saved message. NO red banner about “data policy”, “DLP” or “blocked connector”.", "☐"],
    ["9", ("Run it. ", "Click Test → Manually → Test → Run flow. "
                      "Open Outlook (or the Approvals section in Power Automate) and click Approve on the approval."),
     "The run finishes with a green tick for every step. You received the “Test OK” email, the file "
     "readiness-test.txt is in OneDrive, and a Teams message from Flow bot (if step 7 was used).", "☐"],
    ["10", ("Microsoft Forms. ", "Go to forms.office.com and create a form with 1 question. Back in Power Automate: "
                               "+ Create → Automated cloud flow → trigger “When a new response is submitted” → Create. "
                               "Open the Form Id drop-down."),
     "You can see and select your form in the drop-down. (No need to save this flow.)", "☐"],
    ["11", ("Scheduled flow. ", "+ Create → Scheduled cloud flow. Set it to repeat every 1 day → Create → add any action "
                               "(e.g. Send an email to yourself) → Save → then turn the flow Off."),
     "The flow saves with no error.", "☐"],
], [0.8, 7.2, 7.2, 1.8])

para("When finished: you may delete the test flows (My flows → … → Delete).", italic=True, color=GREY, size=9.5)

# ---------- Results ----------
heading("3. Your result")
t = doc.add_table(rows=3, cols=2)
t.style = "Table Grid"
results = [
    ("C6EFCE", "GREEN — Ready to go",
     "Every step ticked (step 7 may be N/A). Nothing else to do — see you at the training!"),
    ("FFEB9C", "AMBER — Mostly ready",
     "Only step 7 (Teams) failed or is blocked. You can still join — the Teams labs will use an Outlook-email "
     "version instead. Please still send the screenshot to your coordinator so IT is aware."),
    ("FFC7CE", "RED — Need IT help before training",
     "Any other step failed (1–6 or 8–11). Send the screenshot(s) to your training coordinator / IT now — "
     "these features are used in the hands-on labs and must work before Day 1."),
]
for i, (fill, title, text) in enumerate(results):
    a, b = t.rows[i].cells
    shade(a, fill)
    a.text = ""
    a.paragraphs[0].add_run(title).bold = True
    b.text = text
fix_widths(t, [4.5, 12.5])
doc.add_paragraph()

# ---------- For IT ----------
heading("4. For your IT / Microsoft 365 administrator")
para("If a participant gets a RED or AMBER result, please check the following. Only standard Microsoft 365 "
     "features are required — no add-ins, no Copilot, no premium licence.")
table(["Item", "What is needed"], [
    ["Licence", "A Microsoft 365 / Office 365 licence that includes Power Automate for Office 365 "
                "(included in most Business and E3/E5 plans). The Power Automate service plan must be switched ON for the user."],
    ["Environment", "Users can create cloud flows in the Default environment (or a training environment that IT provides)."],
    ["Connectors allowed (DLP policy)",
     "Office 365 Outlook, OneDrive for Business, Excel Online (Business), Approvals, Microsoft Forms, "
     "Microsoft Teams (optional), Office 365 Users (optional), Notifications. "
     "These must be allowed together in the same DLP group (all in “Business”)."],
    ["Approvals", "The Approvals feature can be used in the environment (it is set up automatically the first time "
                  "someone uses it — this may require the Default environment to have Dataverse provisioned)."],
    ["Browser access", "make.powerautomate.com, forms.office.com, OneDrive and Outlook on the web are not blocked "
                       "by proxy / firewall on the training laptop and on the training-room Wi-Fi."],
], [4.5, 12.5])

callout("Stuck? Contact your trainer", [
    ("Mohd Arif Mohd Ghazali", " — AI & Digital Transformation Consultant"),
    "Send your screenshots and this sheet before the training day so we can sort it out in advance.",
], "EAF1FB")

doc.save(OUT)
print("saved", OUT)
